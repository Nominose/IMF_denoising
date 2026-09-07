"""Aggregate tmp/eval_brain.py's per-case CSV into the brain-CT results workbook.

The workbook has one sheet per NFE (NFE=2 ... NFE=50) with rows DDIM / EDM / flow matching /
improved mean flow / imf+GAN and columns K=10 (MAE, SSIM, LPIPS) then K=20 (MAE, SSIM, LPIPS), each
cell "mean+-std" (3 decimals; LPIPS 4). Cells are written with the style of the DDIM cell in the same
column, so the sheet keeps its look (等线 12, thin borders, no bold, no merges).

    python tmp/fill_brain_xlsx.py --rows EDM FM                    # fill only the EDM / flow matching rows
    python tmp/fill_brain_xlsx.py --rows EDM FM --patients common  # same, restricted to the patients every method has
    python tmp/fill_brain_xlsx.py --rows all --patients common --out ..._n14common.xlsx

--check prints, for the rows NOT being written, the CSV aggregate next to what the sheet already
holds, which is how the evaluator is validated against the earlier numbers.
"""
import argparse
import csv
import os
import shutil
from collections import defaultdict
from copy import copy

import numpy as np
from openpyxl import load_workbook

B = r'D:\research\projects\denoising\results\brianct'
ROW_LABEL = {'DDIM': 'DDIM', 'EDM': 'EDM', 'FM': 'flow matching', 'iMF': 'improved mean flow', 'GAN': 'imf+GAN'}
COLS = [('B', 10, 'mae', 3), ('C', 10, 'ssim', 3), ('D', 10, 'lpips', 4),
        ('E', 20, 'mae', 3), ('F', 20, 'ssim', 3), ('G', 20, 'lpips', 4)]


def load_csv(path):
    d = defaultdict(dict)                                    # (method, nfe, k) -> {patient: (mae, ssim, lpips)}
    with open(path, newline='') as f:
        for r in csv.DictReader(f):
            if r['error']:
                continue
            d[(r['method'], int(r['nfe']), int(r['k']))][int(r['patient'])] = (float(r['mae']), float(r['ssim']), float(r['lpips']))
    return d


def fmt(vals, dec):
    v = np.asarray(vals, dtype=float)
    return f'{v.mean():.{dec}f}+-{v.std(ddof=1):.{dec}f}'


def row_of(ws, label):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value and str(row[0].value).strip().lower() == label.lower():
            return row[0].row
    raise KeyError(label)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--csv', default=os.path.join(B, 'brainCT_metrics_per_case.csv'))
    ap.add_argument('--xlsx', default=os.path.join(B, 'results_collection_brainCT.xlsx'))
    ap.add_argument('--out', default=None, help='default: write the workbook in place (a *_before_fill backup is kept)')
    ap.add_argument('--rows', nargs='+', default=['EDM', 'FM'], help="method keys to write, or 'all'")
    ap.add_argument('--patients', choices=['all', 'common'], default='all',
                    help="'common' = only patients present for EVERY method at that NFE and K")
    ap.add_argument('--check', action='store_true', help='compare CSV aggregates of the untouched rows with the sheet')
    a = ap.parse_args()
    rows = list(ROW_LABEL) if a.rows == ['all'] else a.rows
    data = load_csv(a.csv)
    out = a.out or a.xlsx
    if out == a.xlsx:
        shutil.copyfile(a.xlsx, a.xlsx.replace('.xlsx', '_before_fill_backup.xlsx'))
    wb = load_workbook(a.xlsx)
    for ws in wb.worksheets:
        nfe = int(ws.title.split('=')[1])
        ref_row = row_of(ws, 'DDIM')
        for col, k, metric, dec in COLS:
            idx = {'mae': 0, 'ssim': 1, 'lpips': 2}[metric]
            if a.patients == 'common':
                sets = [set(data[(m, nfe, k)]) for m in ROW_LABEL if (m, nfe, k) in data]
                common = set.intersection(*sets) if sets else set()
            for m in ROW_LABEL:
                cell = ws[f'{col}{row_of(ws, ROW_LABEL[m])}']
                per = data.get((m, nfe, k), {})
                if a.patients == 'common':
                    per = {p: v for p, v in per.items() if p in common}
                if m in rows:
                    if not per:
                        cell.value = None; continue
                    cell.value = fmt([v[idx] for v in per.values()], dec)
                    src = ws[f'{col}{ref_row}']
                    cell.font, cell.border, cell.alignment, cell.number_format = copy(src.font), copy(src.border), copy(src.alignment), src.number_format
                elif a.check and per and cell.value:
                    mine = fmt([v[idx] for v in per.values()], dec)
                    flag = '' if mine == str(cell.value).strip() else '   <-- DIFFERS'
                    print(f'  NFE={nfe:<2} {ROW_LABEL[m]:<20} K{k} {metric:<5} sheet {str(cell.value):<16} csv {mine:<16} n={len(per)}{flag}')
    wb.save(out)
    print('wrote', out, '| rows:', [ROW_LABEL[m] for m in rows], '| patients:', a.patients)


if __name__ == '__main__':
    main()
