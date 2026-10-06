r"""Score Mayo EDM / flow-matching baselines stored in the flat layout and fill the Mayo workbook.

    results/mayo/EDM_mayo/NFE{n}/<pid>/pred_img_scans{10,20}.nii.gz
    results/mayo/flowmatching_mayo/NFE{n}/<pid>/pred_img_scans{10,20}.nii.gz
    results/mayo/gt_and_conditions/<pid>/random_0/gt_img.nii.gz          (slices 150-200 of the scan)

Same convention as tmp/eval_mayo.py: abdomen window [-160, 240] HU, mask = voxels whose ORIGINAL GT
is in the window, prediction and GT clipped to the window before MAE / SSIM (data_range 400) / LPIPS
(AlexNet, [-1, 1]); per-slice values averaged over the 50 slices, then mean +- std (ddof=1) over the
3 test patients. Some flow-matching outputs hold the full 240-slice scan; those are cropped to
150-200 first (indexing them like the 50-slice crops lands on the pelvis).

    python tmp/eval_mayo_flat.py --nfe 20                      # print
    python tmp/eval_mayo_flat.py --nfe 20 --check 10 30        # also re-score NFE 10/30 and compare with the sheet
    python tmp/eval_mayo_flat.py --nfe 20 --fill               # write the EDM / flow matching cells of sheet NFE=20
"""
import argparse
import os
import shutil
import sys
from copy import copy

import numpy as np
import nibabel as nb

sys.stdout.reconfigure(encoding='utf-8')
M = os.path.join(os.environ.get('RESULTS_ROOT', r'D:\research\projects\denoising\results'), 'mayo')
XLSX = os.path.join(M, 'results_collection_mayo.xlsx')
VMIN, VMAX = -160.0, 240.0
PATIENTS = ['L192', 'L291', 'L310']
FOLDER = {'EDM': 'EDM_mayo', 'flow matching': 'flowmatching_mayo'}
_fn = None


def metrics(a, g):
    global _fn
    import torch
    from skimage.metrics import structural_similarity
    if _fn is None:
        import lpips
        _fn = lpips.LPIPS(net='alex', verbose=False)
    mae, ssim, lp = [], [], []
    for s in range(g.shape[-1]):
        gs = g[:, :, s]; m = ((gs >= VMIN) & (gs <= VMAX)).astype(np.float32); den = m.sum()
        if den == 0:
            continue
        x, y = np.clip(a[:, :, s], VMIN, VMAX), np.clip(gs, VMIN, VMAX)
        mae.append(float((np.abs(x - y) * m).sum() / den))
        _, smap = structural_similarity(x, y, data_range=VMAX - VMIN, full=True)
        ssim.append(float((smap * m).sum() / den))
        xn = (x - VMIN) / (VMAX - VMIN) * 2 - 1; yn = (y - VMIN) / (VMAX - VMIN) * 2 - 1
        with torch.no_grad():
            lp.append(float(_fn(torch.from_numpy(np.stack([xn] * 3)[None]), torch.from_numpy(np.stack([yn] * 3)[None])).item()))
    assert len(mae) >= 10, 'too few usable slices'
    return float(np.mean(mae)), float(np.mean(ssim)), float(np.mean(lp))


def score(method, nfe, k):
    """-> array (n_patients, 3) or None if any patient is missing"""
    rows = []
    for pid in PATIENTS:
        p = os.path.join(M, FOLDER[method], f'NFE{nfe}', pid, f'pred_img_scans{k}.nii.gz')
        if not os.path.isfile(p):
            return None
        a = np.asarray(nb.load(p).dataobj, np.float32)
        g = np.asarray(nb.load(os.path.join(M, 'gt_and_conditions', pid, 'random_0', 'gt_img.nii.gz')).dataobj, np.float32)
        if a.shape[2] > g.shape[2]:
            a = a[:, :, 150:200]                       # full scan -> the evaluated sub-stack
        rows.append(metrics(a, g))
    return np.array(rows)


def cells(r):
    f = lambda v, d: f'{v.mean():.{d}f} ± {v.std(ddof=1):.{d}f}'
    return [f(r[:, 0], 3), f(r[:, 1], 3), f(r[:, 2], 4)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--nfe', type=int, nargs='+', required=True)
    ap.add_argument('--check', type=int, nargs='*', default=[], help='NFEs to re-score and compare with the sheet')
    ap.add_argument('--fill', action='store_true')
    a = ap.parse_args()
    from openpyxl import load_workbook
    wb = load_workbook(XLSX)

    def row_of(ws, label):
        for row in ws.iter_rows(min_col=1, max_col=1):
            if row[0].value and str(row[0].value).strip().lower() == label.lower():
                return row[0].row
        raise KeyError(label)

    for nfe in a.check:
        ws = wb[f'NFE={nfe}']
        for method in FOLDER:
            r10, r20 = score(method, nfe, 10), score(method, nfe, 20)
            if r10 is None or r20 is None:
                continue
            mine = cells(r10) + cells(r20)
            sheet = [str(ws.cell(row_of(ws, method), c).value) for c in range(2, 8)]
            print(f'[check] NFE={nfe:<2} {method:<14} ' + ('MATCH' if mine == sheet else 'DIFFERS') + f'\n         sheet {sheet}\n         mine  {mine}', flush=True)

    changed = False
    for nfe in a.nfe:
        ws = wb[f'NFE={nfe}']
        ref = row_of(ws, 'DDIM')
        for method in FOLDER:
            r10, r20 = score(method, nfe, 10), score(method, nfe, 20)
            if r10 is None or r20 is None:
                print(f'NFE={nfe} {method}: no data'); continue
            vals = cells(r10) + cells(r20)
            print(f'NFE={nfe:<2} {method:<14} K10 {vals[:3]}  K20 {vals[3:]}   per-patient K20 MAE {np.round(r20[:, 0], 2).tolist()}', flush=True)
            if a.fill:
                for c, v in zip(range(2, 8), vals):
                    cell, src = ws.cell(row_of(ws, method), c), ws.cell(ref, c)
                    cell.value = v
                    cell.font, cell.border, cell.alignment, cell.number_format = copy(src.font), copy(src.border), copy(src.alignment), src.number_format
                changed = True
    if changed:
        shutil.copyfile(XLSX, XLSX.replace('.xlsx', '_before_fill_backup.xlsx'))
        wb.save(XLSX); print('wrote', XLSX)


if __name__ == '__main__':
    main()
